"""Unit detection and metadata extraction (spec sections 15, 16)."""
from __future__ import annotations

from typing import Any, Dict, List, Optional

import pandas as pd

from .core import AuditContext, Dataset
from .static_data import UNIT_TOKENS
from .utils import example_values


def _detect_unit_in_name(name: str) -> Optional[tuple]:
    n = str(name).lower()
    # sort by token length desc so "megawatt-hour" wins over "watt"
    hits = []
    for token, canonical, kind in UNIT_TOKENS:
        t = token.lower()
        if t in n:
            hits.append((len(t), canonical, kind, token))
    if not hits:
        return None
    hits.sort(reverse=True)
    return hits[0][1], hits[0][2], hits[0][3]


def build_unit_catalog(ds: Dataset) -> pd.DataFrame:
    rows = []
    if ds.is_scientific:
        info = ds.scientific_info
        for var in info.get("data_vars", []):
            vd = info.get("variables", {}).get(var, {})
            unit = vd.get("units") or vd.get("attrs", {}).get("units")
            long_name = vd.get("long_name") or vd.get("attrs", {}).get("long_name")
            rows.append({
                "dataset": ds.display_name,
                "variable": var,
                "unit": unit if unit else "UNIT_UNKNOWN",
                "evidence": "NetCDF variable attribute 'units'" if unit else f"long_name={long_name}; no units attribute",
                "confidence": "HIGH" if unit else "MEDIUM",
                "source": "netcdf_attrs",
            })
        return pd.DataFrame(rows)

    if not ds.is_tabular or ds.df is None:
        return pd.DataFrame()

    df = ds.df
    # unit columns (separate columns declaring the unit of another column)
    unit_cols = [c for c in df.columns if str(c).lower().strip() in
                 ("unit", "units", "uom", "unit_of_measure", "measurement_unit")]
    unit_lookup = {}
    if unit_cols:
        ucol = unit_cols[0]
        try:
            unit_lookup = {str(k): str(v) for k, v in
                           df.groupby(ucol, dropna=True).size().to_dict().items()} if False else {}
            # build mapping: for each variable-like column pair is too ambiguous; instead
            # record which units appear in the unit column
            unit_lookup = {str(v): str(v) for v in df[ucol].dropna().unique()}
        except Exception:
            unit_lookup = {}

    for col in df.columns:
        unit = None
        evidence = ""
        confidence = "UNKNOWN"
        src = "none"
        hit = _detect_unit_in_name(str(col))
        if hit:
            unit, kind, token = hit
            evidence = f"column name contains '{token}'"
            confidence = "MEDIUM"
            src = "column_name"
        elif str(col) in unit_lookup:
            unit = str(col)
            evidence = "value referenced in a unit column"
            confidence = "MEDIUM"
            src = "unit_column"
        if unit is None and len(unit_cols) and str(col) not in unit_cols:
            # can't reliably pair values to units without more structure
            unit = "UNIT_UNKNOWN"
            evidence = f"unit column '{unit_cols[0]}' present but pairing ambiguous"
            confidence = "LOW"
            src = "unit_column_ambiguous"
        rows.append({
            "dataset": ds.display_name,
            "variable": str(col),
            "unit": unit if unit else "UNIT_UNKNOWN",
            "evidence": evidence if evidence else "no unit signal in name/metadata",
            "confidence": confidence,
            "source": src,
        })
    return pd.DataFrame(rows)


def extract_excel_metadata(path: str) -> dict:
    """Excel-specific metadata: sheet names, dims, hidden state (spec 16/17)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = {"sheets": []}
        for ws in wb.worksheets:
            out["sheets"].append({
                "name": ws.title,
                "state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
            })
        out["n_sheets"] = len(wb.sheetnames)
        out["sheet_names"] = wb.sheetnames
        return out
    except Exception as exc:
        return {"error": str(exc)}


def run_units_metadata_audit(ctx: AuditContext) -> None:
    unit_rows = []
    meta_rows = []
    for ds in ctx.datasets:
        if not ds.readable:
            continue
        # units
        uc = build_unit_catalog(ds)
        ds.unit_catalog = uc
        if uc is not None and len(uc):
            unit_rows.append(uc)
        # metadata
        meta = dict(ds.metadata)
        if ds.is_scientific:
            info = ds.scientific_info
            meta.update({
                "dims": info.get("dims"),
                "coords": info.get("coords"),
                "data_vars": info.get("data_vars"),
                "attrs": info.get("attrs"),
                "n_variables": len(info.get("data_vars", [])),
            })
        if ds.file_type in ("XLSX", "XLS", "XLSM"):
            meta["excel"] = extract_excel_metadata(ds.path)
        meta_rows.append({
            "dataset": ds.display_name,
            "file_id": ds.file_id,
            "file_type": ds.file_type,
            "metadata_keys": ", ".join(sorted(str(k) for k in meta.keys())[:40]),
            "metadata_json": _jsonish(meta),
        })
    ctx.add_table("08_unit_catalog.csv", pd.concat(unit_rows, ignore_index=True) if unit_rows else pd.DataFrame())
    ctx.add_table("09_metadata_report.csv", pd.DataFrame(meta_rows))


def _jsonish(obj: Any, maxlen: int = 2000) -> str:
    import json
    try:
        s = json.dumps(obj, default=str)
        return s[:maxlen]
    except Exception:
        return str(obj)[:maxlen]
